"""Generates small, non-sensitive document fixtures in memory.

Not committed as binary files (especially DOCX, a zip archive) — generated
on the fly instead, so every fixture's exact content is reviewable as plain
Python rather than an opaque blob in the repo. Not a test module itself (no
`test_*` functions), so pytest won't collect it.
"""

from __future__ import annotations

from io import BytesIO


def make_txt_bytes() -> bytes:
    return (
        b"Hello World.\n\n"
        b"This is a test TXT file for AegisAI document ingestion.\n"
        b"It has multiple lines.\n\n\n\n"
        b"And some extra blank lines above that should be normalized."
    )


def make_markdown_bytes() -> bytes:
    return b"# Sample Markdown\n\nThis is **markdown** content.\n\n- item one\n- item two\n"


def make_docx_bytes() -> bytes:
    from docx import Document as DocxDocument

    document = DocxDocument()
    document.add_paragraph("Hello DOCX World.")
    document.add_paragraph("This is a second paragraph for testing extraction.")
    document.add_paragraph("")
    document.add_paragraph("A third paragraph after a blank one.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_empty_docx_bytes() -> bytes:
    from docx import Document as DocxDocument

    buffer = BytesIO()
    DocxDocument().save(buffer)
    return buffer.getvalue()


def make_pdf_bytes(*, text: str = "Hello PDF World") -> bytes:
    """A minimal, hand-built single-page PDF with one text-drawing
    instruction — small enough to be legible inline here, and sufficient
    for pypdf's text extraction. The xref table is deliberately not fully
    accurate; pypdf tolerates and repairs this (verified against the
    installed pypdf version), which real-world "slightly malformed" PDFs
    also routinely require.
    """
    content_stream = f"BT /F1 18 Tf 10 100 Td ({text}) Tj ET".encode()
    return b"\n".join(
        [
            b"%PDF-1.4",
            b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj",
            b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj",
            b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]"
            b"/Resources<</Font<</F1 4 0 R>>>>/Contents 5 0 R>>endobj",
            b"4 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj",
            b"5 0 obj<</Length " + str(len(content_stream)).encode() + b">>",
            b"stream",
            content_stream,
            b"endstream",
            b"endobj",
            b"xref",
            b"0 6",
            b"0000000000 65535 f ",
            b"trailer<</Size 6/Root 1 0 R>>",
            b"startxref",
            b"0",
            b"%%EOF",
        ]
    )


def make_corrupt_pdf_bytes() -> bytes:
    """Has the right magic bytes (passes the sniff check) but no valid
    structure after that — for exercising the extraction-failure path."""
    return b"%PDF-1.4\nthis is not a real pdf body at all, just noise\n%%EOF"
